"""Date-scoped, permission-filtered Frappe prepared communication reports."""
from datetime import date, datetime, time, timedelta
import frappe
from relay_core.business_inbox import messages, check_type
from relay_core.permissions import require_access

REPORT = 'Relay Communication History'
MAX_ROWS = 20000


def date_range(from_date, to_date):
    try:
        if not isinstance(from_date,str) or not isinstance(to_date,str):raise ValueError
        start=date.fromisoformat(from_date);end=date.fromisoformat(to_date)
        if start.isoformat()!=from_date or end.isoformat()!=to_date or end<start:raise ValueError
        return datetime.combine(start,time.min),datetime.combine(end+timedelta(days=1),time.min)
    except (ValueError,OverflowError):
        frappe.throw('Choose a valid start and end date; the end date must not precede the start date')


def validated_filters(filters):
    require_access()
    values=frappe.parse_json(filters) if isinstance(filters,str) else filters
    if not isinstance(values,dict) or set(values)-{'business_doctype','business_name','business_names','from_date','to_date','output_format','include_images','report_mode'}:
        frappe.throw('Invalid communication report filters')
    date_range(values.get('from_date'),values.get('to_date'))
    doctype=values.get('business_doctype')
    check_type(doctype)
    names=values['business_names'] if 'business_names' in values else [values.get('business_name')]
    if not isinstance(names,list) or not 1<=len(names)<=20 or any(not isinstance(name,str) or not name or len(name)>140 for name in names):
        frappe.throw('Choose between 1 and 20 business records')
    if len(set(names))!=len(names):frappe.throw('Choose each business record only once')
    for name in names:frappe.get_doc(doctype,name).check_permission('read')
    include_images=values.get('include_images',0)
    if type(include_images) not in (bool,int) or include_images not in (0,1):frappe.throw('Invalid image option')
    mode=values.get('report_mode','Full messages')
    if mode not in ('Full messages','Daily summary'):frappe.throw('Choose full messages or daily summary')
    output=values.get('output_format','PDF')
    if output not in ('PDF','Excel'):frappe.throw('Choose PDF or Excel')
    return {**{key:values[key] for key in ('business_doctype','from_date','to_date')},'business_name':names[0],'business_names':names,'output_format':output,'include_images':int(include_images),'report_mode':mode}


def source_report(filters=None):
    values=validated_filters(filters)
    result=[];seen={}
    for business_name in values['business_names']:
        cursor=None
        while True:
            page=messages(values['business_doctype'],business_name,before=cursor,limit=100,
                          from_date=values['from_date'],to_date=values['to_date'])
            from relay_core.report_media import inventory
            for row in page['items']:
                if row['name'] in seen:
                    seen[row['name']]['business_records'] += '; '+business_name
                    continue
                result.append({'relay_contact':row['relay_contact'],'business_records':business_name,'message':row['name'],'contact':row['contact_label'],
                    'channel':row['channel'],'direction':row['direction'],'status':row['status'],
                    'message_time':row.get('sent_at') or row['creation'],
                    'message_type':row.get('message_type') or 'text','content':row.get('content') or '',
                    'attachments':'; '.join(item['type']+': '+item['name'] for item in inventory(row))})
                seen[row['name']]=result[-1]
            if len(result)>MAX_ROWS:
                frappe.throw('This report exceeds 20,000 accessible messages. Choose a shorter date range; no partial report was generated.')
            if not page['has_more']:break
            next_cursor=page['first']
            if not next_cursor or next_cursor==cursor:frappe.throw('Report paging could not progress')
            cursor=next_cursor
    result.sort(key=lambda row:(str(row['message_time']),row['message']))
    columns=[{'fieldname':key,'label':label,'fieldtype':kind,'width':width} for key,label,kind,width in (
        ('business_records','Business records','Data',240),('message_time','Time','Datetime',160),('contact','Contact','Data',180),
        ('channel','Channel','Data',160),('direction','Direction','Data',90),
        ('status','Status','Data',90),('message_type','Type','Data',90),
        ('content','Message','Small Text',400),('attachments','Attachments','Small Text',240),('message','Message reference','Data',180))]
    return columns,result


def execute(filters=None):
    values=validated_filters(filters)
    columns,rows=source_report(values)
    if values['report_mode']=='Daily summary':
        from relay_core.report_summaries import summarize
        rows=summarize(rows)
        columns = [column for column in columns if column['fieldname'] != 'message'] + [
            {'fieldname':'summary','label':'Summary reference','fieldtype':'Data','width':180},
            {'fieldname':'source_count','label':'Source message count','fieldtype':'Int','width':140},
            {'fieldname':'source_references','label':'Source message references','fieldtype':'Small Text','width':400},
        ]
    return columns,rows


def source_names(rows):
    return {name for row in rows for name in row.get('sources',[row['message']])}


@frappe.whitelist(methods=['POST'])
def request(filters):
    values=validated_filters(filters)
    from frappe.core.doctype.prepared_report.prepared_report import make_prepared_report
    return make_prepared_report(REPORT,filters=values)


@frappe.whitelist(methods=['GET'])
def status(name):
    require_access()
    doc=frappe.get_doc('Prepared Report',name)
    if doc.report_name!=REPORT or doc.owner!=frappe.session.user:
        frappe.throw('This report belongs to another user',frappe.PermissionError)
    # Relay's API authorizes its own report owner and every source record.
    # Do not require the unrelated Desk Prepared Report User role here.
    if not permission(doc):
        frappe.throw('Access to this report changed. Prepare a new report.',frappe.PermissionError)
    values=validated_filters(doc.filters)
    return {'name':doc.name,'status':doc.status,'created':doc.creation,'filters':values,
            'finished':doc.report_end_time,'ready':doc.status=='Completed' and artifact_path(doc).is_file(),'error':'Report generation failed. Review the job in Desk.' if doc.status=='Error' else ''}


def notify(doc, method=None):
    if doc.report_name == REPORT and doc.status in ('Completed','Error'):
        if doc.status=='Completed':
            try:build_artifact(doc)
            except Exception:
                doc.db_set('status','Error',update_modified=False)
                frappe.log_error(title='Relay report export failed')
        frappe.publish_realtime('relay_changed',{'kind':'reports'},user=doc.owner,after_commit=True)


def artifact_path(doc):
    import hashlib
    from pathlib import Path
    extension='pdf' if validated_filters(doc.filters)['output_format']=='PDF' else 'xlsx'
    return Path(frappe.get_site_path('private','relay_reports')) / (hashlib.sha256(doc.name.encode()).hexdigest()+'.'+extension)


def prepared_data(doc):
    import json
    data=json.loads(doc.get_prepared_data())
    rows=data.get('result')
    if not isinstance(rows,list) or len(rows)>MAX_ROWS or any(not isinstance(row,dict) for row in rows):
        frappe.throw('Prepared report data is invalid')
    return data


def build_artifact(doc):
    from io import BytesIO
    import os
    from html import escape
    values=validated_filters(doc.filters)
    data=prepared_data(doc);rows=data['result'];columns=data['columns']
    from relay_core.report_media import collect
    previews,media_notes=collect(rows,bool(values['include_images']))
    if values['output_format']=='Excel':
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        workbook=Workbook(write_only=True);sheet=workbook.create_sheet('Daily summaries' if values['report_mode']=='Daily summary' else 'Messages')
        def literal(value):
            cell=WriteOnlyCell(sheet,value=str(value or ''));cell.data_type='s';return cell
        sheet.append([literal(col['label']) for col in columns])
        for row in rows:sheet.append([literal(row.get(col['fieldname'])) for col in columns])
        if previews or media_notes:
            from openpyxl.drawing.image import Image as SheetImage
            media_sheet=workbook.create_sheet('Images')
            media_sheet.append(['Message reference','Preview / availability'])
            position=2
            for row in rows:
                name=row['message']
                if name not in previews and name not in media_notes:continue
                media_sheet.append([name,media_notes.get(name,'Image preview')])
                if name in previews:
                    image=SheetImage(BytesIO(previews[name]));scale=min(180/image.width,180/image.height);image.width*=scale;image.height*=scale
                    media_sheet.add_image(image,'B'+str(position))
                    for _ in range(12):media_sheet.append(['',''])
                    position+=12
                position+=1
        output=BytesIO();workbook.save(output);content=output.getvalue()
    else:
        from weasyprint import HTML
        esc=lambda value:escape(str(value or ''))
        image_resources={'relay-image:'+name:raw for name,raw in previews.items()}
        def deny(url,*args,**kwargs):
            if url in image_resources:return {'string':image_resources[url],'mime_type':'image/jpeg'}
            raise ValueError('External report resources are not allowed')
        blocks=[]
        for row in rows:
            blocks.append('<article><header>'+esc(row.get('business_records'))+' · '+esc(row['message_time'])+' · '+esc(row['contact'])+' · '+esc(row['channel'])+' · '+esc(row['direction'])+' · '+esc(row['status'])+'</header><div>'+esc(row['content'])+'</div><small>'+esc(row.get('source_references') or row['message'])+' · '+esc(row['message_type'])+'</small><p>'+esc(row.get('attachments'))+'</p>'+('<img style="max-width:100%;max-height:350px" src="'+esc('relay-image:'+row['message'])+'">' if row['message'] in previews else '<p>'+esc(media_notes.get(row['message'],''))+'</p>')+'</article>')
        count_label = (str(len(rows))+' daily summaries · '+str(len(source_names(rows)))+' source messages') if values['report_mode']=='Daily summary' else str(len(rows))+' messages'
        html='<html><meta charset="utf-8"><style>@page{size:A4;margin:16mm;@bottom-right{content:counter(page)}}body{font-family:sans-serif;font-size:10pt}article{padding:12px 0;border-bottom:1px solid #ddd}header,small{color:#555;font-size:8pt}article div{white-space:pre-wrap;overflow-wrap:anywhere;margin:6px 0}h1{font-size:18pt}</style><h1>Communication report</h1><p>'+esc(values['business_doctype'])+' · '+esc('; '.join(values['business_names']))+'</p><p>'+esc(values['from_date'])+' through '+esc(values['to_date'])+' · '+esc(count_label)+'</p>'+''.join(blocks)+'</html>'
        content=HTML(string=html,url_fetcher=deny).write_pdf()
    path=artifact_path(doc);path.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
    temporary=path.with_suffix(path.suffix+'.tmp')
    with open(temporary,'wb') as stream:
        os.chmod(temporary,0o600);stream.write(content)
    temporary.replace(path)


@frappe.whitelist(methods=['GET'])
def download(name):
    from werkzeug.wrappers import Response
    state=status(name)
    if not state['ready']:frappe.throw('Report is not ready for download')
    doc=frappe.get_doc('Prepared Report',name)
    # status() already rechecked owner and every stored source against current
    # permissions. Do not repeat that potentially 20,000-message scan here.
    path=artifact_path(doc)
    mime='application/pdf' if path.suffix=='.pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return Response(path.read_bytes(),content_type=mime,headers={'Content-Disposition':'attachment; filename="relay-communication-report'+path.suffix+'"','Cache-Control':'private, no-store','X-Content-Type-Options':'nosniff'})


def cleanup(doc, method=None):
    if doc.report_name!=REPORT:return
    import hashlib
    from pathlib import Path
    root=Path(frappe.get_site_path('private','relay_reports'))
    key=hashlib.sha256(doc.name.encode()).hexdigest()
    for extension in ('.pdf','.xlsx','.pdf.tmp','.xlsx.tmp'):
        (root/(key+extension)).unlink(missing_ok=True)


def permission(doc, user=None, permission_type=None):
    if doc.report_name != REPORT:return True
    user=user or frappe.session.user
    if user!=doc.owner or user!=frappe.session.user:return False
    try:
        validated_filters(doc.filters)
        if doc.status=='Completed':
            saved=prepared_data(doc)
            _,current=source_report(doc.filters)
            allowed={row['message'] for row in current}
            if not source_names(saved['result']).issubset(allowed):return False
    except (frappe.PermissionError,frappe.ValidationError,frappe.DoesNotExistError):
        return False
    return True


def file_permission(doc, user=None, permission_type=None):
    if doc.attached_to_doctype!='Prepared Report' or not doc.attached_to_name:return True
    report=frappe.get_doc('Prepared Report',doc.attached_to_name)
    return permission(report,user,permission_type)


@frappe.whitelist(methods=['GET'])
def recent():
    require_access()
    rows=frappe.get_all('Prepared Report',filters={'report_name':REPORT,'owner':frappe.session.user},
        fields=['name','status','creation','filters'],order_by='creation desc',limit_page_length=20)
    items=[]
    for row in rows:
        try:values=validated_filters(row.filters)
        except (frappe.PermissionError,frappe.ValidationError,frappe.DoesNotExistError):continue
        items.append({'name':row.name,'status':row.status,'created':row.creation,'filters':values})
    return {'items':items}
