"""Recipient rows from a sheet or JSON, WhatsApp template parameters and text or media messages for bulk campaigns."""
import csv
import ipaddress
import mimetypes
import re
from decimal import Decimal, InvalidOperation
from urllib.parse import unquote, urlsplit
import phonenumbers
from relay_core.channels import channel
from relay_core.identity import phone_number

PHONE_COLUMNS = ('phone', 'mobile', 'mobile no', 'mobile number', 'phone number', 'whatsapp', 'number')
MEDIA_FORMATS = ('IMAGE', 'VIDEO', 'DOCUMENT')
COMPONENTS = ('HEADER', 'BODY', 'FOOTER', 'BUTTONS')
BUTTONS = ('QUICK_REPLY', 'URL', 'PHONE_NUMBER', 'COPY_CODE')
MAX_ROWS = 100000
MAX_VALUE = 1024
MESSAGE_MEDIA = {'photo': 'image', 'video': 'video', 'document': 'document'}  # Registry kind -> Meta media type.
PLACEHOLDER = re.compile(r'\{\{\s*([^{}]+?)\s*\}\}')
PRIVATE_SUFFIXES = ('.localhost', '.local', '.internal', '.intranet', '.lan', '.home.arpa')
HEADER_KINDS = {media: kind for kind, media in MESSAGE_MEDIA.items()}  # Meta header type -> registry kind.


def sheet_phone(value, region):
    """Mobile E.164 digits, accepting Excel's exact integer text representations."""
    text = str(value or '').strip()
    if re.fullmatch(r'[0-9]+\.0+', text):
        text = text.split('.')[0]
    elif re.fullmatch(r'[+-]?[0-9]+(?:\.[0-9]+)?[eE][+-]?[0-9]+', text):
        try:
            number = Decimal(text)
        except InvalidOperation:
            return None
        if (len(number.as_tuple().digits) > 15 or not 0 < number < 10**15
                or number != number.to_integral_value()):
            return None
        text = str(int(number))
    destination = phone_number(text, region)
    if destination:
        number = phonenumbers.parse('+' + destination, None)
        if (phonenumbers.is_valid_number(number) and phonenumbers.number_type(number) in
                (phonenumbers.PhoneNumberType.MOBILE, phonenumbers.PhoneNumberType.FIXED_LINE_OR_MOBILE)):
            return destination
    return None


def template_fields(template):
    """Fields a sheet row must supply, in Meta component order. Raises ValueError for parts campaigns cannot fill."""
    named = str(template.get('parameter_format') or 'POSITIONAL').upper() == 'NAMED'
    fields = []
    for component in template.get('components') or []:
        kind = str(component.get('type') or '').upper()
        text = component.get('text') or ''
        header = str(component.get('format') or 'TEXT').upper()
        if kind not in COMPONENTS:
            raise ValueError(f'Campaigns cannot send {kind.lower() or "unknown"} templates')
        if kind == 'HEADER' and header not in (*MEDIA_FORMATS, 'TEXT'):
            raise ValueError(f'Campaigns cannot send templates with a {header.lower()} header')
        if kind == 'HEADER' and header in MEDIA_FORMATS:
            fields.append({'key': 'header_media', 'label': 'Header ' + header.lower() + ' attachment',
                           'component': 'header', 'media': header.lower()})
        elif kind in ('HEADER', 'BODY'):
            for variable in variables(text, named):
                fields.append({'key': f'{kind.lower()}:{variable}', 'label': f'{kind.capitalize()} value {variable}',
                               'component': kind.lower(), 'name': variable if named else None})
        elif kind == 'BUTTONS':
            fields.extend(button_fields(component.get('buttons') or []))
    return fields


def button_fields(buttons):
    fields = []
    for index, button in enumerate(buttons):
        kind = str(button.get('type') or '').upper()
        if kind not in BUTTONS:
            raise ValueError(f"Campaigns cannot send templates with {kind.replace('_', ' ') or 'unknown'} buttons")
        label = button.get('text') or 'Button'
        if kind == 'URL' and variables(button.get('url') or '', False):
            fields.append({'key': f'button:{index}', 'label': f'{label} link value', 'component': 'button', 'index': index})
        elif kind == 'COPY_CODE':
            fields.append({'key': f'button:{index}', 'label': f'{label} coupon code', 'component': 'button',
                           'index': index, 'sub_type': 'copy_code'})
    return fields


def variables(text, named):
    pattern = r'\{\{\s*([a-z][a-z0-9_]*)\s*\}\}' if named else r'\{\{\s*(\d+)\s*\}\}'
    found = []
    for value in re.findall(pattern, text):
        if value not in found:
            found.append(value)
    return found if named else sorted(found, key=int)


def build_components(fields, values, header_media=None):
    grouped, buttons = {}, []
    for field in fields:
        value = values[field['key']]
        if field['component'] == 'button':
            parameter = ({'type': 'coupon_code', 'coupon_code': value} if field.get('sub_type') == 'copy_code'
                         else {'type': 'text', 'text': value})
            buttons.append({'type': 'button', 'sub_type': field.get('sub_type') or 'url', 'index': str(field['index']),
                            'parameters': [parameter]})
            continue
        target = grouped.setdefault(field['component'], {'type': field['component'], 'parameters': []})
        if field.get('media'):
            if isinstance(value, dict) and not header_media:
                raise ValueError('Upload the header attachment before sending')
            target['parameters'].append({'type': field['media'], field['media']: header_media or {'link': value}})
        else:
            parameter = {'type': 'text', 'text': value}
            if field.get('name'):
                parameter['parameter_name'] = field['name']
            target['parameters'].append(parameter)
    return [*grouped.values(), *buttons]


def read_rows(path):
    """Yield (sheet row number, header -> text) without loading the workbook."""
    lower = path.lower()
    if lower.endswith(('.csv', '.tsv')):
        with open(path, newline='', encoding='utf-8-sig') as handle:
            yield from labelled(csv.reader(handle, delimiter='\t' if lower.endswith('.tsv') else ','))
    elif lower.endswith('.xlsx'):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from labelled(book.worksheets[0].iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ValueError('Upload a .xlsx, .csv or .tsv file')


def labelled(source):
    header = None
    for number, raw in enumerate(source, start=1):
        cells = [text(value) for value in raw]
        if header is None:
            header = cells
            if not any(header) or len(set(filter(None, header))) != len([c for c in header if c]):
                raise ValueError('The first row must hold unique column names')
            continue
        if not any(cells):
            continue
        if number - 1 > MAX_ROWS:
            raise ValueError(f'A campaign supports at most {MAX_ROWS} rows')
        yield number, {column: cells[i] if i < len(cells) else '' for i, column in enumerate(header) if column}


def given_rows(rows, after=0):
    """(row number, column -> text) for JSON objects, numbered on from after, the same shape read_rows yields."""
    for number, row in enumerate(rows, start=after + 1):
        if (not isinstance(row, dict) or not all(isinstance(column, str) and column.strip() for column in row)
                or any(isinstance(value, (dict, list)) for value in row.values())):
            raise ValueError(f'Row {number} must map column names to plain values')
        yield number, {column.strip(): text(value) for column, value in row.items()}


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def columns(path, sample=5):
    rows = []
    header = []
    for _, row in read_rows(path):
        header = header or list(row)
        rows.append(row)
        if len(rows) >= sample:
            break
    return header, rows


def suggest(header, fields, *, first_phone=False):
    """Phone column by name; named values by their column name; positional template values take the
    remaining columns left to right."""
    by_name = {column.lower(): column for column in header}
    mapping = {'phone': header[0] if first_phone and header else next((by_name[c] for c in PHONE_COLUMNS if c in by_name), ''), 'fields': {}, 'fixed': {}}
    positional = [field for field in fields if is_positional(field)]
    for field in fields:
        if field in positional:
            continue
        if field['key'] == 'header_media':
            candidates = (header_column(field).lower(), field['media'], field['key'], field['label'].lower(),
                          'image url', 'media url', 'url', 'file', 'file name')
        else:
            candidates = ((field.get('name') or field['key'].split(':')[-1]).lower(), field['key'].lower(),
                          field['label'].lower(), 'image url', 'media url', 'url')
        for candidate in candidates:
            if candidate in by_name:
                mapping['fields'][field['key']] = by_name[candidate]
                break
    used = {mapping['phone'], *mapping['fields'].values()}
    mapping['fields'].update(zip((field['key'] for field in positional), (c for c in header if c not in used)))
    return mapping


def header_column(field):
    """Sheet column name of a template's header media, as the Excel template writes it."""
    return f"Header {field['media']} (file name in the ZIP or https link)"


def is_positional(field):
    return field['component'] in ('header', 'body') and not field.get('media') and not field.get('name')


def check_mapping(mapping, header, fields):
    if not isinstance(mapping, dict) or set(mapping) - {'phone', 'fields', 'fixed', 'extra', 'header_attachment'}:
        raise ValueError('Invalid column mapping')
    from relay_core.campaign_media import check_attachment
    attachment = check_attachment(mapping.get('header_attachment'), fields)
    extra = mapping.get('extra') or []
    if not isinstance(extra, list) or len(extra) > 10 or any(column not in header for column in extra):
        raise ValueError('Choose up to 10 extra columns from the file')
    columns_used, fixed = mapping.get('fields') or {}, mapping.get('fixed') or {}
    if not isinstance(columns_used, dict) or not isinstance(fixed, dict):
        raise ValueError('Invalid column mapping')
    if mapping.get('phone') not in header:
        raise ValueError('Choose the phone number column')
    keys = {field['key'] for field in fields}
    if set(columns_used) - keys or set(fixed) - keys:
        raise ValueError('Mapping names a value this template does not use')
    for key, column in columns_used.items():
        if column not in header:
            raise ValueError(f'Column {column} is not in the file')
    by_number = attachment and attachment['mode'] != 'named'
    if by_number and ('header_media' in columns_used or 'header_media' in fixed):
        raise ValueError('Choose one header attachment source')
    if attachment and not by_number and 'header_media' not in columns_used:
        raise ValueError('Choose the column that names each row\'s header file')
    for field in fields:
        if field['key'] == 'header_media' and by_number:
            continue
        if field['key'] not in columns_used and not str(fixed.get(field['key']) or '').strip():
            raise ValueError(f"Choose a column or fixed value for {field['label']}")
        if field['key'] in fixed and field.get('media') and not (field['component'] == 'media' and is_site_file(fixed[field['key']])):
            media_url(fixed[field['key']])
    return {'phone': mapping['phone'], 'fields': dict(columns_used), 'extra': list(dict.fromkeys(extra)),
            **({'header_attachment': attachment} if attachment else {}),
            'fixed': {key: str(value).strip() for key, value in fixed.items() if key not in columns_used}}


def media_url(value):
    parts = urlsplit(value)
    if parts.scheme != 'https' or not parts.hostname or parts.username or parts.password or parts.fragment or len(value) > 2048:
        raise ValueError('Media needs an HTTPS link without credentials')
    return value


def public_media_url(value):
    """The HTTPS link when its host is public. Meta fetches it, so internal names and addresses are refused."""
    host = urlsplit(media_url(value)).hostname.rstrip('.').lower()
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    last = host.rsplit('.', 1)[-1]
    if address and not address.is_global or not address and (
            '.' not in host or host.endswith(PRIVATE_SUFFIXES) or last.isdigit() or last.startswith('0x')):
        raise ValueError('Media link must point to a public host')
    return value


def zip_name(value):
    """The ZIP file a header cell names, as its last path part; a link names no file. Match it without case."""
    return '' if '://' in value else re.split(r'[\\/]', value.strip())[-1]


def header_value(value, field, zip_files):
    """(header value, problem) of a row's header cell: a public HTTPS link of the header's type, or the
    checked ZIP file it names."""
    if name := zip_name(value):
        return zip_files.get(name.casefold(), (None, f'Header file {name[:120]} is not in the ZIP'))
    try:
        media_url(value)
    except ValueError:
        return None, f"{field['label']} must be an HTTPS link"
    try:
        public_media_url(value)
    except ValueError:
        return None, f"{field['label']} must be a public link"
    mime = mimetypes.guess_type(urlsplit(value).path)[0]
    if mime and mime not in channel('WhatsApp').features['media'][HEADER_KINDS[field['media']]]:
        return None, f"Wrong file type for {field['media']} header"
    return value, None


def is_site_file(value):
    """An uploaded File that every row sends; Relay uploads it to Hub once when the campaign starts."""
    return isinstance(value, str) and value.startswith(('/files/', '/private/files/'))


def prepare_row(row, mapping, fields, region, header_file=None, zip_files=None):
    """Return (destination, values, problem). Problems are user-facing reasons. A number-matched header file
    comes as header_file; a header column names a link or one of zip_files (ZIP name -> (file, problem))."""
    destination = sheet_phone(row.get(mapping['phone']), region)
    if not destination:
        return None, None, 'Invalid phone number'
    values = {}
    by_number = (mapping.get('header_attachment') or {}).get('mode') in ('static', 'dynamic')
    for field in fields:
        if field['key'] == 'header_media' and by_number:
            if not header_file:
                return destination, None, 'Missing header attachment'
            values[field['key']] = header_file
            continue
        column = mapping['fields'].get(field['key'])
        value = row.get(column, '') if column else mapping['fixed'].get(field['key'], '')
        if not value:
            return destination, None, f"Missing {field['label']}"
        if len(value) > MAX_VALUE:
            return destination, None, f"{field['label']} is too long"
        if column and field['key'] == 'header_media':
            value, problem = header_value(value, field, zip_files or {})
            if problem:
                return destination, None, problem
        elif column and field.get('media'):  # check_mapping checked the fixed value.
            try:
                media_url(value)
            except ValueError:
                return destination, None, f"{field['label']} must be an HTTPS link"
        values[field['key']] = value
    extra = {column: row.get(column, '')[:MAX_VALUE] for column in mapping.get('extra', [])}
    return destination, {'values': values, **({'extra': extra} if extra else {})}, None


def message_fields(message):
    """Sheet values a text or media message uses: each {{column}} placeholder, then the media link."""
    fields = [{'key': 'text:' + column, 'label': column, 'component': 'text', 'name': column}
              for column in dict.fromkeys(PLACEHOLDER.findall(message.get('body') or message.get('caption') or ''))]
    if message['kind'] in MESSAGE_MEDIA:
        fields.append({'key': 'media_url', 'label': message['kind'].capitalize() + ' link', 'component': 'media',
                       'media': MESSAGE_MEDIA[message['kind']]})
    return fields


def check_message(message):
    """Normalised text or media message. Raises ValueError with the reason."""
    features = channel('WhatsApp').features
    kind = message.get('kind') if isinstance(message, dict) else None
    if kind == 'text':
        body = str(message.get('body') or '').strip()
        if not body or len(body) > features['max_text_length']:
            raise ValueError(f"Write a message of 1-{features['max_text_length']} characters")
        return {'kind': kind, 'body': body}
    if kind not in MESSAGE_MEDIA or kind not in features['media']:
        raise ValueError('Choose text, photo, video or document')
    caption = str(message.get('caption') or '').strip()
    if len(caption) > features['max_caption_length']:
        raise ValueError(f"A caption holds at most {features['max_caption_length']} characters")
    return {'kind': kind, 'caption': caption}


def accepted_types(kind):
    return channel('WhatsApp').features['media'][kind] if kind in MESSAGE_MEDIA else ()


def fill(text, values):
    return PLACEHOLDER.sub(lambda match: values.get('text:' + match.group(1), ''), text)


def message_problem(message, values):
    """Reason a filled row cannot be sent, else None."""
    features = channel('WhatsApp').features
    if message['kind'] == 'text' and len(fill(message['body'], values)) > features['max_text_length']:
        return 'Message is too long'
    if len(fill(message.get('caption') or '', values)) > features['max_caption_length']:
        return 'Caption is too long'
    return None


def build_message(message, values, media_id=None):
    """Hub messages.send or messages.media arguments for one row, without the recipient. A media_id replaces
    the row's link when the campaign sends one uploaded file."""
    if message['kind'] == 'text':
        return {'operation': 'messages.send', 'body': fill(message['body'], values)}
    source = {'media_id': media_id} if media_id else {'media_url': values['media_url']}
    payload = {'operation': 'messages.media', 'media_type': MESSAGE_MEDIA[message['kind']], **source}
    caption = fill(message['caption'], values)
    if caption:
        payload['caption'] = caption
    if message['kind'] == 'document':
        payload['filename'] = file_name(values['media_url'])
    return payload


def file_name(url):
    name = re.sub(r'[\\/\r\n]', '', unquote(urlsplit(url).path.rsplit('/', 1)[-1]))
    return name[:255] or 'document'
