"""Local DB/API import and send checks; no live providers, migrations or cleanup deletes.

Run with FRAPPE_STREAM_LOGGING=1 ../../env/bin/python -B tests/integration_campaign_sheet_site.py.
Files use repository-local fixtures through the file adapter, keeping disk writes in this repo.
Clearly named campaigns and their contacts remain on the development test site for inspection.
"""
import csv
import json
import os
import secrets
import zipfile
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import frappe
from frappe import _dict
from relay_core import bulk_campaigns as bulk, bulk_sheet, bulk_updates, campaign_media, test_mode, whatsapp, settings, inbound
from relay_core.automation.service import internal_save
from relay_core.contact_lifecycle import endpoint_key

ROOT = Path(__file__).resolve().parents[1]
SITES = ROOT.parents[1] / 'sites'
SITE = 'relay-flows-test.localhost'
os.chdir(ROOT / '.runtime')
frappe.init(site=SITE, sites_path=str(SITES))
assert frappe.conf.db_host in (None, 'localhost', '127.0.0.1')
frappe.connect(); frappe.set_user('Administrator'); frappe.flags.mute_emails = True
assert 'relay_core' in frappe.get_installed_apps()
assert 'sheet_row' in frappe.db.get_table_columns(bulk.RECIPIENT)
RUN = 'Local sheet check ' + secrets.token_hex(4)
folder = ROOT / '.runtime' / RUN.replace(' ', '-')
folder.mkdir()
account = RUN.replace(' ', '-')
channel = frappe.get_doc({'doctype': 'Relay Channel', 'channel': 'WhatsApp', 'account': account, 'enabled': 1}).insert()
frappe.db.commit()
# Reserved NANP 555 numbers, accepted as possible E.164 identities, never sent to a real provider.
numbers = [str(12025550100 + index) for index in range(60)]
files = {}


def fixture(name, raw):
    path = folder / name; path.write_bytes(raw)
    url = '/private/files/' + name
    files[url] = _dict(file_name=name, file_url=url, is_private=1, get_full_path=lambda: str(path))
    return {'file_url': url, 'filename': name}


def sheet(name, values):
    path = folder / name
    with path.open('w') as handle:
        writer = csv.writer(handle); writer.writerow(['Recipients', 'Offer'])
        writer.writerows(values)
    return fixture(name, path.read_bytes())['file_url']


class Hub:
    def __init__(self):
        self.uploads, self.messages = [], []
    def execute(self, payload):
        assert payload['operation'] == 'messages.media_upload'
        self.uploads.append(payload)
        return {'id': str(5000 + len(self.uploads))}
    def call_tool(self, tool, payload):
        assert tool == 'whatsapp.bulk'
        self.messages.extend(payload['messages'])
        for row in payload['messages']:
            assert row['message']['components'][0]['parameters'][0]['image']['id'].isdigit()
        return {'deliveries': [{'key': row['key'], 'state': 'queued', 'id': str(7000 + i)}
                              for i, row in enumerate(payload['messages'])], 'queued': 0}


def campaign(label, source, attachment):
    template = {'name': 'local_offer', 'language': 'en', 'components': [
        {'type': 'HEADER', 'format': 'IMAGE'}, {'type': 'BODY', 'text': 'Your offer: {{1}}'}]}
    doc = internal_save(frappe.get_doc({'doctype': 'Relay Campaign', 'title': RUN + ' ' + label,
        'channel': channel.name, 'operation': 'templates.send', 'template_name': template['name'],
        'language': 'en', 'arguments': json.dumps(bulk.template_snapshot(template)), 'status': 'Draft',
        'account': account, 'country': 'US', 'reply_window_hours': 72, 'request_key': account + '-' + label}), new=True)
    inspected = bulk.inspect(doc.name, source)
    assert inspected['mapping']['phone'] == 'Recipients'
    bulk.import_rows(doc.name, {'phone': 'Offer', 'fields': {'body:1': 'Offer'}, 'header_attachment': attachment})
    frappe.db.commit()
    bulk.import_job(doc.name)  # The same background-job entry point the API queued.
    result = bulk.read(doc.name)
    assert not result['last_error'], result
    return frappe.get_doc('Relay Campaign', doc.name), result


try:
    with ExitStack() as stack:
        stack.enter_context(patch.object(test_mode, 'is_enabled', return_value=True))
        stack.enter_context(patch.object(frappe, 'enqueue'))
        original_setting = settings.value
        stack.enter_context(patch.object(settings, 'value', side_effect=lambda key: account if key == 'relay_hub_account' else original_setting(key)))
        stack.enter_context(patch.object(bulk, 'uploaded_file', side_effect=lambda url: files[url]))
        hub = Hub()
        stack.enter_context(patch.object(whatsapp, 'client', return_value=hub))
        stack.enter_context(patch.object(bulk, 'hub_arguments', side_effect=lambda channel, args: args))
        # Existing opted-out number must remain opted out throughout import and sending.
        bulk.sheet_contact(channel.name, numbers[0])
        endpoint = frappe.db.get_value('Relay Contact Endpoint', {'endpoint_key': endpoint_key(channel.name, numbers[0])}, 'name')
        frappe.db.set_value('Relay Contact Endpoint', endpoint, 'marketing_opt_out', 1)
        frappe.db.commit()
        image = fixture('offer.png', b'\x89PNG\r\n\x1a\nlocal-only-header')
        source = sheet('static.csv', [(number, '25 percent') for number in numbers] + [(numbers[1], '25 percent'), ('bad', 'Invalid')])
        doc, result = campaign('static', source, {'mode': 'static', 'files': [image]})
        assert result['contact_counts'] == {'created': 59, 'existing': 2, 'invalid': 1}, result['contact_counts']
        assert frappe.db.count(bulk.RECIPIENT, {'campaign': doc.name, 'status': 'Pending'}) == 59
        assert frappe.db.get_value('Relay Contact Endpoint', endpoint, 'marketing_opt_out') == 1
        doc.status = 'Running'; internal_save(doc); frappe.db.commit()
        bulk.send_job(doc.name)
        assert len(hub.uploads) == 1 and len(hub.messages) == 59, (len(hub.uploads), len(hub.messages))
        assert frappe.db.count('Relay Message', {'source_key': ['like', 'bulk-%'], 'channel': channel.name}) == 59
        # Dynamic preview: exact match, missing, duplicate, wrong type, national match.
        dynamic = [fixture(numbers[1] + '.png', b'\x89PNG\r\n\x1a\n'), fixture(numbers[3] + '.png', b'\x89PNG\r\n\x1a\n'),
                   fixture(numbers[3] + '.jpg', b'\xff\xd8\xff'), fixture(numbers[4] + '.pdf', b'%PDF-1.4'),
                   fixture(numbers[5][1:] + '.png', b'\x89PNG\r\n\x1a\n')]
        source = sheet('dynamic.csv', [(number, '50 percent') for number in numbers[1:6]])
        doc, result = campaign('dynamic', source, {'mode': 'dynamic', 'files': dynamic})
        assert result['contact_counts'] == {'created': 0, 'existing': 5, 'invalid': 0}
        assert sorted((row['reason'], row['count']) for row in result['issues']) == [
            ('Duplicate header attachments', 1), ('Missing header attachment', 1), ('Wrong file type for image header', 1)]
        doc.status = 'Running'; internal_save(doc); frappe.db.commit()
        bulk.send_job(doc.name)
        assert len(hub.uploads) == 3 and len(hub.messages) == 61
        bulk.send_job(doc.name)
        assert len(hub.uploads) == 3 and len(hub.messages) == 61
        # A ZIP takes the same sheet validation and send path; File bytes stay in this repository.
        bundle = BytesIO()
        with zipfile.ZipFile(bundle, 'w') as archive:
            column = bulk_sheet.header_column({'media': 'image'})
            archive.writestr('sheet/people.csv', f'Recipients,{column},Offer\n{numbers[6]},offer.png,25 percent\n'
                             f'{numbers[7]},missing.png,50 percent\n{numbers[8]},media/OFFER.PNG,75 percent\nbad,offer.png,Invalid')
            archive.writestr('media/offer.png', b'\x89PNG\r\n\x1a\nzip-offer')
            archive.writestr('__MACOSX/._people.csv', 'ignored')
        source = fixture('people.zip', bundle.getvalue())['file_url']
        template = {'name': 'local_offer', 'language': 'en', 'components': [
            {'type': 'HEADER', 'format': 'IMAGE'}, {'type': 'BODY', 'text': '{{1}}'}]}
        zipdoc = internal_save(frappe.get_doc({'doctype': 'Relay Campaign', 'title': RUN + ' ZIP',
            'channel': channel.name, 'operation': 'templates.send', 'template_name': 'local_offer', 'language': 'en',
            'arguments': json.dumps(bulk.template_snapshot(template)), 'status': 'Draft', 'account': account,
            'country': 'US', 'reply_window_hours': 72, 'request_key': account + '-zip'}), new=True)
        def save_zip_file(filename, raw, campaign):
            return _dict(file_url=fixture(filename, raw)['file_url'])
        with patch.object(campaign_media, 'save_private', side_effect=save_zip_file):
            inspected = bulk.inspect(zipdoc.name, source)
        bulk.import_rows(zipdoc.name, inspected['mapping'])
        frappe.db.commit()
        before = len(hub.messages)
        bulk.import_job(zipdoc.name)
        assert len(hub.messages) == before  # Validate never calls Hub.
        result = bulk.read(zipdoc.name)
        assert result['validated'] and len(result['error_rows']) == 2, result
        assert frappe.db.count(bulk.RECIPIENT, {'campaign': zipdoc.name, 'status': 'Pending'}) == 2
        bulk.error_report(zipdoc.name)
        from openpyxl import load_workbook
        report = load_workbook(BytesIO(frappe.local.response.filecontent))
        assert report.active.max_row == 3
        report.close()
        with patch.object(whatsapp, 'checked_client', return_value=(hub, {})), patch.object(bulk, 'approved_template', return_value=template):
            bulk.control(zipdoc.name, 'start')
        frappe.db.commit()
        uploads = len(hub.uploads)
        bulk.send_job(zipdoc.name)
        assert len(hub.messages) == before + 2 and len(hub.uploads) == uploads + 1  # Two rows reuse one ZIP file.
        recipients = frappe.get_all(bulk.RECIPIENT, filters={'campaign': zipdoc.name, 'status': 'Queued'},
            fields=['name', 'request_key', 'arguments'], order_by='sheet_row')
        assert all(json.loads(row.arguments)['validation']['status'] == 'Validated' for row in recipients)
        event = inbound.insert_event({'account': account, 'event_key': account + '-zip-result',
            'hub_event_id': account + '-zip-result', 'payload': '{}'}, inline=True)
        with patch('relay_core.bulk_updates.connection_value', return_value=account):
            bulk_updates.project_batch(event, [{'key': row.request_key, 'id': 7000 + index, 'state': 'submitted',
                'status': 'sent' if index == 0 else 'failed', 'message_id': 'local-only-zip-result-' + str(index),
                'error': {} if index == 0 else {'code': 131026, 'title': 'Local-only simulated failure'}}
                for index, row in enumerate(recipients)])
        frappe.db.commit()
        assert bulk_updates.recount(zipdoc.name)['status'] == 'Completed'
        assert frappe.db.get_value(bulk.RECIPIENT, recipients[0].name, 'status') == 'Sent'
        assert frappe.db.get_value(bulk.RECIPIENT, recipients[1].name, 'status') == 'Failed'
        assert frappe.db.get_value(bulk.RECIPIENT, recipients[1].name, 'reason') == 'Meta error 131026: Local-only simulated failure'
        assert [entry['phase'] for entry in bulk.read(zipdoc.name)['phases']] == ['Uploaded', 'Validating', 'Validated', 'Running', 'Completed']
        # Separate DB connections exercise the real Redis lock and unique endpoint index.
        concurrent_number = '12025550160'
        def enroll(_):
            frappe.init(site=SITE, sites_path=str(SITES)); frappe.connect(); frappe.set_user('Administrator')
            frappe.flags.mute_emails = True
            try:
                return bulk.sheet_contact(channel.name, concurrent_number)
            finally:
                frappe.db.rollback(); frappe.destroy()
        with ThreadPoolExecutor(max_workers=4) as pool:
            outcomes = list(pool.map(enroll, range(4)))
        assert outcomes.count(('created', None)) == 1 and outcomes.count(('existing', None)) == 3, outcomes
        print(json.dumps({'site': SITE, 'fixture': RUN, 'static_rows': 62, 'dynamic_rows': 5,
            'concurrent_imports': 4, 'messages_admitted': len(hub.messages), 'provider_uploads': len(hub.uploads), 'result': 'passed'}))
finally:
    frappe.db.rollback()
    frappe.destroy()
