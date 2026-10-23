"""Native prepared report lifecycle on the isolated Frappe database."""
import json, os
from datetime import date
from pathlib import Path
from unittest.mock import patch
root=Path(__file__).resolve().parents[1];sites=root/'.runtime/automation-bench/sites'
assert Path(os.environ['RELAY_AUTOMATION_TEST_SITES']).resolve()==sites.resolve()
os.chdir(sites)
import frappe
from relay_core import prepared_reports
from frappe.core.doctype.prepared_report import prepared_report as framework
frappe.init(site='relay-ai-test.local')
assert frappe.conf.db_name=='relay_ai_test_20260922' and frappe.conf.db_port==18433
frappe.connect();frappe.set_user('Administrator');frappe.flags.mute_emails=True
frappe.reload_doc('relay_core','report','relay_communication_history',force=True)
frappe.clear_cache()
business=frappe.get_all('Relay Channel Simulation Business',filters={'title':['like','Types business %']},pluck='name',order_by='creation desc',limit_page_length=1)[0]
values={'business_doctype':'Relay Channel Simulation Business','business_name':business,'from_date':'1970-01-01','to_date':'2100-01-01','include_images':1}
job=None;extra_business=None;linked_contact=None;dated_message=None
from io import BytesIO
from PIL import Image
from relay_core import report_media
image_bytes=BytesIO();Image.new('RGB',(20,10),'blue').save(image_bytes,format='JPEG')
media_fixture=patch.object(report_media,'image_preview',return_value=image_bytes.getvalue())
media_fixture.start()
try:
 columns,rows=prepared_reports.execute(values)
 assert rows and all(row['message'] for row in rows)
 # Own the date boundary instead of relying on an earlier script's timestamp mix.
 dated_message=(rows[0]['message'],frappe.db.get_value('Relay Message',rows[0]['message'],'sent_at'))
 frappe.db.set_value('Relay Message',dated_message[0],'sent_at','2000-01-01 12:00:00',update_modified=False)
 columns,rows=prepared_reports.execute(values)
 extra_business=frappe.get_doc({'doctype':'Relay Channel Simulation Business','title':'Prepared report shared scope'}).insert()
 relay_contact=frappe.get_doc('Relay Contact',frappe.db.get_value('Relay Message',rows[0]['message'],'relay_contact'))
 linked_contact=frappe.get_doc('Contact',relay_contact.contact)
 linked_contact.append('links',{'link_doctype':extra_business.doctype,'link_name':extra_business.name});linked_contact.save()
 values['business_names']=[business,extra_business.name]
 _,shared=prepared_reports.execute(values)
 assert len(shared)==len(rows),'Shared contact messages duplicated across business scopes'
 assert any(extra_business.name in row['business_records'] and business in row['business_records'] for row in shared)
 for names in ([],[business,business],[business]*21):
  try:prepared_reports.validated_filters(dict(values,business_names=names))
  except frappe.ValidationError:pass
  else:raise AssertionError('Invalid multi-record selection accepted')

 today=date.today().isoformat()
 _,bounded=prepared_reports.execute(dict(values,from_date=today,to_date=today))
 assert all(str(row['message_time']).startswith(today) for row in bounded)
 assert len(bounded)<len(rows),'Fixture must include older provider timestamps'
 for changes in ({'from_date':'2026-02-30'},{'from_date':'2026-09-25','to_date':'2026-09-24'},{'to_date':None}):
  try:prepared_reports.validated_filters(dict(values,**changes))
  except frappe.ValidationError:pass
  else:raise AssertionError('Invalid report date accepted')
 with patch.object(framework,'enqueue') as enqueue:
  job=prepared_reports.request(values)['name']
  assert enqueue.call_count==1
 assert prepared_reports.status(job)['status']=='Queued'
 frappe.db.commit()
 with patch.object(frappe,'publish_realtime') as published:
  framework.generate_report(job)
  assert any(call.args[:2]==('relay_changed',{'kind':'reports'}) and call.kwargs.get('user')=='Administrator' and call.kwargs.get('after_commit') for call in published.call_args_list), [(call.args,call.kwargs) for call in published.call_args_list]
 frappe.db.commit()
 assert prepared_reports.status(job)['status']=='Completed'
 assert prepared_reports.status(job)['ready']
 assert prepared_reports.permission(frappe.get_doc('Prepared Report',job)) is True
 assert prepared_reports.file_permission(frappe._dict(attached_to_doctype='Relay Message',attached_to_name='fixture')) is True
 assert prepared_reports.permission(frappe._dict(report_name='Another Report')) is True
 pdf=prepared_reports.download(job).get_data()
 assert pdf.startswith(b'%PDF-') and b'/Subtype /Image' in pdf
 data=json.loads(frappe.get_doc('Prepared Report',job).get_prepared_data())
 assert len(data['result'])==len(rows)
 artifact=prepared_reports.artifact_path(frappe.get_doc('Prepared Report',job))
 frappe.delete_doc('Prepared Report',job,ignore_permissions=True);frappe.db.commit()
 assert not artifact.exists(),'Deleted prepared report left an export behind'
 with patch.object(framework,'enqueue'):
  job=prepared_reports.request(dict(values,output_format='Excel'))['name']
 frappe.db.commit();framework.generate_report(job);frappe.db.commit()
 assert prepared_reports.status(job)['ready']
 from io import BytesIO
 from openpyxl import load_workbook
 raw=prepared_reports.download(job).get_data()
 import zipfile
 with zipfile.ZipFile(BytesIO(raw)) as archive:assert any(name.startswith('xl/media/') for name in archive.namelist())
 sheet=load_workbook(BytesIO(raw),read_only=True).active
 assert sheet.max_row is None or sheet.max_row==len(rows)+1
 spreadsheet=list(sheet.rows)
 assert len(spreadsheet)==len(rows)+1 and all(cell.data_type!='f' for row in spreadsheet for cell in row)
 frappe.set_user('channels-executor@example.invalid')
 try:prepared_reports.status(job)
 except frappe.PermissionError:pass
 else:raise AssertionError('Report ownership not enforced')
 print('PASS native prepared report queued/completed/private data, date filtering, invalid dates and owner isolation')
finally:
 media_fixture.stop()
 frappe.set_user('Administrator')
 if dated_message:frappe.db.set_value('Relay Message',dated_message[0],'sent_at',dated_message[1],update_modified=False)
 if job:frappe.delete_doc('Prepared Report',job,ignore_permissions=True)
 if linked_contact and extra_business:
  linked_contact.reload();linked_contact.set('links',[row for row in linked_contact.links if row.link_name!=extra_business.name]);linked_contact.save()
 if extra_business:frappe.delete_doc(extra_business.doctype,extra_business.name,ignore_permissions=True)
 frappe.db.commit();frappe.destroy()
