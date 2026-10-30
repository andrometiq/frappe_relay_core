"""ZIP and three-phase campaign checks. All persistence and Hub calls are local doubles."""
import json
import stat
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import frappe
from frappe import _dict
from openpyxl import load_workbook
from pypdf.errors import PdfReadError
from relay_core import bulk_campaigns as bulk, bulk_sheet, bulk_updates, campaign_media
from site_file_limit import use_relay_file_limit

NOW = '2026-10-30 12:00:00'
USER = 'local-campaign-manager@example.invalid'


def setUpModule():
    use_relay_file_limit()


class ZipTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory()))

    def archive(self, entries):
        path = self.root / 'campaign.zip'
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, raw in entries:
                archive.writestr(name, raw)
        return str(path)

    def test_nested_files_are_flattened_and_hidden_files_ignored(self):
        path = self.archive([('sheet/people.csv', 'Mobile,Name\n12025550101,Asha'),
            ('media/12025550101.png', b'\x89PNG\r\n\x1a\n'), ('__MACOSX/extra.csv', 'ignore'),
            ('.DS_Store', 'ignore'), ('media/.hidden.txt', 'ignore'), ('.hidden/extra.xlsx', 'ignore')])
        sheet, files = campaign_media.zip_contents(path)
        self.assertEqual(sheet, 'people.csv')
        self.assertEqual([name for name, _ in files], ['people.csv', '12025550101.png'])

    def test_path_traversal_absolute_windows_paths_and_links_are_rejected(self):
        for name in ('../escape.png', 'media/../../escape.png', '/escape.png', 'C:\\escape.png', '..\\escape.png'):
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, 'unsafe'):
                campaign_media.zip_contents(self.archive([('people.csv', 'Mobile'), (name, b'x')]))
        link = zipfile.ZipInfo('link.png')
        link.create_system = 3
        link.external_attr = (stat.S_IFLNK | 0o777) << 16
        with self.assertRaisesRegex(ValueError, 'symbolic link'):
            campaign_media.zip_contents(self.archive([('people.csv', 'Mobile'), (link, '/etc/passwd')]))

    def test_missing_or_two_sheets_are_rejected(self):
        for entries in ([], [('a.png', b'x')], [('a.csv', 'Mobile'), ('b.xlsx', b'x')]):
            with self.subTest(entries=entries), self.assertRaisesRegex(ValueError, 'exactly one'):
                campaign_media.zip_contents(self.archive(entries))

    def test_expanded_size_limit_prevents_zip_bomb(self):
        path = self.archive([('people.csv', 'x' * 8192)])
        self.assertLess(Path(path).stat().st_size, 1024)
        with patch.object(campaign_media, 'ZIP_MAX_BYTES', 1024), self.assertRaisesRegex(ValueError, '200 MiB'):
            campaign_media.zip_contents(path)

    def test_archive_size_and_file_count_are_bounded(self):
        path = self.archive([('a.csv', 'Mobile'), ('b.png', b'x')])
        with patch.object(campaign_media, 'ZIP_MAX_FILES', 1), self.assertRaisesRegex(ValueError, '10000'):
            campaign_media.zip_contents(path)
        with patch.object(campaign_media, 'ZIP_MAX_BYTES', 1), self.assertRaisesRegex(ValueError, '200 MiB'):
            campaign_media.zip_contents(path)

    def test_duplicate_flattened_names_and_unsupported_extensions(self):
        for entries, reason in [([('a/people.csv', 'x'), ('b/PEOPLE.CSV', 'x')], 'duplicate'),
                ([('people.csv', 'x'), ('run.exe', 'x')], 'only a sheet'),
                ([('people.csv', 'x'), ('nested.zip', 'x')], 'only a sheet')]:
            with self.subTest(entries=entries), self.assertRaisesRegex(ValueError, reason):
                campaign_media.zip_contents(self.archive(entries))

    def test_corrupt_zip_has_friendly_reason(self):
        path = self.root / 'bad.zip'; path.write_bytes(b'not a zip')
        with self.assertRaisesRegex(ValueError, 'could not be read'):
            campaign_media.zip_contents(str(path))

    def test_inspect_uses_existing_private_header_mapping_and_invalidates_previous_rows(self):
        path = self.archive([('people.csv', 'Mobile,Name\n12025550101,Asha'), ('12025550101.png', b'\x89PNG\r\n\x1a\n')])
        fields = bulk_sheet.template_fields({'components': [{'type': 'HEADER', 'format': 'IMAGE'}, {'type': 'BODY', 'text': '{{1}}'}]})
        doc = _dict(name='Local ZIP', arguments=json.dumps({'fields': fields, 'validated': True}))
        def save(name, raw, *args, **kwargs):
            self.assertEqual(kwargs['is_private'], 1)
            (self.root / name).write_bytes(raw)
            return _dict(file_url='/private/files/' + name)
        fake = SimpleNamespace(utils=SimpleNamespace(now=lambda: NOW), session=_dict(user=USER))
        with patch.object(bulk, 'frappe', fake), patch.object(bulk, 'bulk', return_value=doc), \
                patch.object(bulk, 'uploaded_file', return_value=_dict(get_full_path=lambda: path)), \
                patch.object(bulk, 'source_path', side_effect=lambda url: str(self.root / url.split('/')[-1])), \
                patch('frappe.utils.file_manager.save_file', side_effect=save), patch.object(bulk, 'internal_save'), \
                patch.object(bulk.whatsapp, 'client') as hub:
            result = bulk.inspect(doc.name, '/private/files/campaign.zip')
        self.assertEqual(result['mapping']['fields'], {'body:1': 'Name'})
        self.assertEqual(result['mapping']['header_attachment']['files'][0]['filename'], '12025550101.png')
        self.assertEqual(result['mapping']['header_attachment']['mode'], 'dynamic')
        self.assertFalse(json.loads(doc.arguments)['validated'])
        self.assertEqual(doc.source_file, '/private/files/people.csv')
        hub.assert_not_called()

    def test_large_or_damaged_header_files_flag_only_their_rows(self):
        path = self.archive([('people.csv', 'Mobile'), ('12025550101.pdf', b'%PDF-ok'),
                             ('12025550102.pdf', b'%PDF-' + b'x' * 3000), ('12025550103.pdf', b'%PDF-broken')])
        def save(name, raw, *args, **kwargs):
            if name == '12025550103.pdf':
                raise PdfReadError('startxref not found')
            return _dict(file_url='/private/files/' + name)
        with patch.object(campaign_media, 'get_media_max_bytes', return_value=2048), \
                patch('frappe.utils.file_manager.save_file', side_effect=save):
            sheet_url, value = campaign_media.unpack(path, 'Local ZIP', 'document')
        self.assertEqual(sheet_url, '/private/files/people.csv')
        self.assertEqual(value['files'], [{'file_url': '/private/files/12025550101.pdf', 'filename': '12025550101.pdf'}])
        self.assertEqual(value['rejected'], [
            {'filename': '12025550102.pdf', 'reason': 'Header attachment must be at most 2 KB'},
            {'filename': '12025550103.pdf', 'reason': 'Wrong file content for document header'}])
        fields = bulk_sheet.template_fields({'components': [{'type': 'HEADER', 'format': 'DOCUMENT'}]})
        self.assertIs(campaign_media.check_attachment(value, fields), value)
        issues = []
        with patch.object(campaign_media, 'checked_file'):
            result = campaign_media.resolve(value, 'document', {'12025550101', '12025550102', '12025550103'}, 'US', issues)
        self.assertEqual(result['12025550101'], (value['files'][0], None))
        self.assertEqual(result['12025550102'], (None, 'Header attachment must be at most 2 KB'))
        self.assertEqual(result['12025550103'], (None, 'Wrong file content for document header'))
        self.assertEqual([issue['filename'] for issue in issues], ['12025550102.pdf', '12025550103.pdf'])

    def test_rejected_files_keep_our_reasons_and_dynamic_mode(self):
        fields = bulk_sheet.template_fields({'components': [{'type': 'HEADER', 'format': 'IMAGE'}]})
        only = {'mode': 'dynamic', 'files': [], 'rejected': [{'filename': 'a.png', 'reason': 'Wrong file content for image header'}]}
        self.assertIs(campaign_media.check_attachment(only, fields), only)
        for rejected, mode in [([{'filename': 'a.png', 'reason': 'Traceback: secret'}], 'dynamic'),
                               ([{'filename': 'a.png', 'reason': ''}], 'dynamic'),
                               ([{'filename': 'a.png', 'reason': 'Header attachment is empty'}], 'static')]:
            with self.subTest(rejected=rejected, mode=mode), self.assertRaises(ValueError):
                campaign_media.check_attachment({'mode': mode, 'files': [] if mode == 'dynamic' else
                    [{'file_url': '/private/files/b.png', 'filename': 'b.png'}], 'rejected': rejected}, fields)


class InspectErrorTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory()))
        self.fake = SimpleNamespace(throw=Mock(side_effect=lambda message: (_ for _ in ()).throw(frappe.ValidationError(message))),
                                    log_error=Mock())
        self.enterContext(patch.object(bulk, 'frappe', self.fake))
        self.enterContext(patch.object(bulk, 'bulk', return_value=_dict(name='Local', arguments=json.dumps({'fields': []}))))

    def test_wrong_file_type_keeps_its_own_message(self):
        with self.assertRaisesRegex(frappe.ValidationError, r'^Upload a \.xlsx, \.csv or \.tsv file$'):
            bulk.inspect('Local', '/private/files/notes.txt')
        self.fake.log_error.assert_not_called()

    def test_damaged_sheet_has_friendly_message(self):
        broken = BytesIO()
        with zipfile.ZipFile(broken, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<broken')
        for name, raw in (('people.xlsx', b'not a workbook'), ('parts.xlsx', broken.getvalue()),
                          ('people.csv', b'Mobile\n\xff\xfe\xfa')):
            path = self.root / name
            path.write_bytes(raw)
            with self.subTest(name=name), patch.object(bulk, 'uploaded_file', return_value=_dict(get_full_path=lambda: str(path))), \
                    self.assertRaisesRegex(frappe.ValidationError, 'sheet could not be read'):
                bulk.inspect('Local', '/private/files/' + name)


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.db = Mock()
        self.fake = SimpleNamespace(db=self.db, utils=SimpleNamespace(now=lambda: NOW), session=_dict(user=USER),
                                    throw=Mock(side_effect=lambda message: (_ for _ in ()).throw(frappe.ValidationError(message))))
        self.enterContext(patch.object(bulk, 'frappe', self.fake))
        self.enterContext(patch.object(bulk, 'sheet_contact', return_value=('existing', None)))
        self.doc = _dict(name='Local validation', channel='WA', country='US', source_file='/private/files/people.csv',
            arguments=json.dumps({'fields': bulk_sheet.template_fields({'components': [{'type': 'BODY', 'text': '{{1}}'}]})}),
            column_map=json.dumps({'phone': 'Mobile', 'fields': {'body:1': 'Name'}, 'fixed': {}}))

    def test_row_errors_and_validated_evidence_persist_without_hub_calls(self):
        rows = [(2, {'Mobile': '12025550101', 'Name': 'Asha'}), (3, {'Mobile': '+1 202 555 0101', 'Name': 'Asha'}),
                (4, {'Mobile': 'bad', 'Name': 'Invalid'}), (5, {'Mobile': '12025550102', 'Name': ''}),
                (6, {'Mobile': '12025550103', 'Name': 'Opted out'}), (7, {'Mobile': '12025550104', 'Name': 'Revoked'})]
        def facts(channel, numbers, blocked=None):
            blocked.add('12025550104')
            return {}, {'12025550103'}
        with patch.object(bulk, 'endpoint_facts', side_effect=facts), patch.object(bulk.whatsapp, 'client') as hub:
            bulk.add_rows(self.doc, rows, set())
        records = self.db.bulk_insert.call_args.args[2]
        self.assertEqual([r[14] for r in records], ['', 'Duplicate row', 'Invalid phone number',
            'Missing Body value 1', 'Opted out', 'Endpoint verification revoked'])
        self.assertEqual([r[10] for r in records], ['Pending'] + ['Skipped'] * 5)
        for record in records:
            audit = json.loads(record[12])['validation']
            self.assertEqual((audit['at'], audit['user'], audit['reason']), (NOW, USER, record[14]))
        self.assertEqual(json.loads(records[0][12])['validation']['status'], 'Validated')
        hub.assert_not_called()

    def test_text_reply_window_checked_during_validation(self):
        self.doc.arguments = json.dumps({'message': {'kind': 'text', 'body': 'Hello'}, 'fields': []})
        with patch.object(bulk, 'endpoint_facts', return_value=({}, set())), \
                patch.object(bulk.service_window, 'open_numbers', return_value={'12025550101'}):
            bulk.insert_chunk(self.doc, [(1, '12025550101', {'values': {}}, None), (2, '12025550102', {'values': {}}, None)])
        self.assertEqual([r[14] for r in self.db.bulk_insert.call_args.args[2]], ['', bulk.OUTSIDE_WINDOW])

    def test_replacement_upload_cannot_send_stale_pending_rows(self):
        self.doc.arguments = json.dumps({'fields': [], 'validated': False})
        self.doc.status = 'Draft'
        self.db.exists.return_value = True
        with self.assertRaisesRegex(frappe.ValidationError, 'Validate'):
            bulk.check_validation(self.doc)

    def test_phase_history_preserves_template_and_records_actor_time(self):
        original = json.loads(self.doc.arguments)['fields']
        bulk.phase(self.doc, 'Validating', validated=False)
        bulk.phase(self.doc, 'Validated', validated=True)
        bulk.phase(self.doc, 'Running')
        snapshot = json.loads(self.doc.arguments)
        self.assertEqual(snapshot['fields'], original)
        self.assertTrue(snapshot['validated'])
        self.assertEqual(snapshot['phases'], [{'phase': step, 'at': NOW, 'user': USER} for step in ('Validating', 'Validated', 'Running')])

    def test_timeout_and_long_queue_use_current_message_count(self):
        self.assertEqual([bulk.send_timeout(n) for n in (0, 1, 120, 121, 100000)], [600, 600, 600, 605, 500000])
        self.fake.enqueue = Mock()
        self.db.count.return_value = 1000
        bulk.enqueue('send_job', self.doc.name)
        self.assertEqual(self.fake.enqueue.call_args.kwargs['queue'], 'long')
        self.assertEqual(self.fake.enqueue.call_args.kwargs['timeout'], 5000)
        self.assertLess(bulk.SLICE_SECONDS + bulk.LOCK_WAIT + 3 * 30, bulk.send_timeout(1))

    def test_cancellation_defers_pending_chunks_at_the_slice_deadline(self):
        with patch.object(bulk.time, 'monotonic', return_value=50), patch.object(bulk.whatsapp, 'client') as hub:
            self.assertTrue(bulk.cancel_remaining(self.doc, deadline=45))
        self.db.sql.assert_not_called()
        hub.assert_not_called()

    def test_unknown_failures_are_not_mistaken_for_friendly_field_errors(self):
        for text in ('Missing secret provider response', 'Header database_password=secret', 'Traceback: secret'):
            self.assertNotIn('secret', bulk.friendly_reason(text))
        fields = json.loads(self.doc.arguments)['fields']
        self.assertEqual(bulk.friendly_reason('Missing Body value 1', fields), 'Missing Body value 1')

    def test_new_download_and_upload_apis_deny_before_reading_files(self):
        with patch.object(bulk, 'require_access', side_effect=frappe.PermissionError), patch.object(bulk, 'uploaded_file') as file:
            for call in (lambda: bulk.sample_sheet('denied'), lambda: bulk.error_report('denied'),
                         lambda: bulk.inspect('denied', '/private/files/a.zip'), lambda: bulk.import_rows('denied', {})):
                with self.assertRaises(frappe.PermissionError):
                    call()
        file.assert_not_called()

    def test_error_report_is_xlsx_and_uploaded_text_cannot_become_formula(self):
        self.fake.local = SimpleNamespace(response={})
        bulk.xlsx_response([['Row', 'Mobile', 'Reason'], [2, '=1+1', 'Missing header attachment']], 'errors.xlsx')
        book = load_workbook(BytesIO(self.fake.local.response['filecontent']))
        self.assertEqual(book.active['B2'].value, '=1+1')
        self.assertEqual(book.active['B2'].data_type, 's')
        book.close()

    def test_excel_template_has_phone_first_and_reuses_shared_named_values(self):
        self.fake.local = SimpleNamespace(response={})
        self.doc.arguments = json.dumps({'fields': [
            {'key': 'header_media', 'label': 'Header image attachment'},
            {'key': 'header:name', 'name': 'name', 'label': 'Header value name'},
            {'key': 'body:name', 'name': 'name', 'label': 'Body value name'}]})
        with patch.object(bulk, 'bulk', return_value=self.doc):
            bulk.sample_sheet(self.doc.name)
        book = load_workbook(BytesIO(self.fake.local.response['filecontent']))
        self.assertEqual(list(book.active.values), [('mobile number', 'name')])
        book.close()

    def test_error_report_pages_all_errors_and_hides_raw_failures(self):
        self.fake.get_all = Mock(side_effect=[[_dict(sheet_row=2, destination='12025550101', status='Skipped', reason='Missing header attachment')],
            [_dict(sheet_row=90, destination='12025550102', status='Failed', reason='Traceback: secret provider payload')], []])
        with patch.object(bulk, 'bulk', return_value=self.doc), patch.object(bulk, 'xlsx_response') as response:
            bulk.error_report(self.doc.name)
            rows = list(response.call_args.args[0])
        self.assertEqual([r[0] for r in rows], ['Row', 2, 90])
        self.assertNotIn('secret', rows[2][3])
        self.assertEqual(self.fake.get_all.call_args.kwargs['filters']['sheet_row'], ['>', 90])

    def test_completion_adds_phase_without_losing_validation_history(self):
        bulk.phase(self.doc, 'Validated', validated=True)
        fake = SimpleNamespace(db=Mock(sql=Mock(return_value=[('Sent', 1, 0), ('Failed', 1, 0)]),
                                      get_value=Mock(return_value='Running')), get_doc=Mock(return_value=self.doc),
                               utils=SimpleNamespace(now_datetime=lambda: NOW))
        with patch.object(bulk_updates, 'frappe', fake):
            result = bulk_updates.recount(self.doc.name)
        self.assertEqual(result['status'], 'Completed')
        saved = fake.db.set_value.call_args.args[2]
        self.assertEqual([entry['phase'] for entry in json.loads(saved['arguments'])['phases']], ['Validated', 'Completed'])
